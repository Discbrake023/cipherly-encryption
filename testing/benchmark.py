"""
testing/benchmark.py
=====================
Skrip pengujian menyeluruh untuk Tugas Proyek Aplikasi Kriptografi - Topik A.

Menjalankan seluruh "Pengujian wajib" pada spesifikasi tugas:
  1. Kebenaran dekripsi pada >= 10 masukan berbeda (termasuk citra & PDF)
  2. Waktu enkripsi/dekripsi untuk berkas 1 KB, 1 MB, 10 MB
  3. Avalanche effect (perubahan 1 bit plaintext, perubahan 1 bit kunci)
  4. Entropi & histogram byte, ciphertext vs plaintext
  5. Perbandingan AES-256-GCM vs ChaCha20-Poly1305

Keluaran:
  laporan/hasil_pengujian.xlsx   -> seluruh tabel hasil (dikumpulkan sebagai
                                     "Data pengujian" sesuai Bagian 5)
  laporan/histogram_plaintext.png
  laporan/histogram_ciphertext.png

Cara menjalankan (dari root proyek):
    python -m testing.benchmark
"""

from __future__ import annotations

import os
import statistics
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import crypto_core as ck
from testing.metrics import (
    avalanche_effect_bit_flip,
    avalanche_effect_key_change,
    shannon_entropy,
    byte_histogram,
)
from testing.sample_files import make_png, make_pdf

ALGORITHMS = ["aes-gcm", "chacha20"]
KDF_FOR_SPEED_TESTS = "pbkdf2"   # KDF paling cepat untuk pengujian throughput enkripsi
PASSWORD = "K@t4S4ndiUjiCoba2026!"
TIMING_REPEATS = 5               # jumlah pengulangan tiap pengukuran waktu, dirata-rata
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "laporan")


# ---------------------------------------------------------------------------
# 1) Uji korektnas: >= 10 masukan berbeda, termasuk citra & PDF
# ---------------------------------------------------------------------------

def build_test_inputs() -> dict:
    inputs = {
        "teks_pendek": "Halo, ini pesan rahasia singkat.".encode("utf-8"),
        "teks_kosong_1_byte": b"\x00",
        "teks_unicode": "Keamanan Informasi 🔒 - Universitas Siliwangi".encode("utf-8"),
        "angka_biner": bytes(range(256)),
        "data_acak_1kb": os.urandom(1024),
        "data_acak_64kb": os.urandom(64 * 1024),
        "citra_png_sintetis": make_png(64, 64),
        "citra_png_besar": make_png(256, 256),
        "dokumen_pdf_sintetis": make_pdf(),
        "teks_panjang_lorem": ("Lorem ipsum dolor sit amet. " * 500).encode("utf-8"),
        "data_semua_nol": bytes(4096),
        "data_semua_0xff": bytes([0xFF]) * 4096,
    }
    # Jika mahasiswa menaruh berkas asli di sample_data/, ikutkan juga.
    sample_dir = os.path.join(os.path.dirname(OUTPUT_DIR), "sample_data")
    if os.path.isdir(sample_dir):
        for fname in sorted(os.listdir(sample_dir)):
            fpath = os.path.join(sample_dir, fname)
            if os.path.isfile(fpath):
                with open(fpath, "rb") as fh:
                    inputs[f"sample_data/{fname}"] = fh.read()
    return inputs


def run_correctness_tests(inputs: dict) -> list[dict]:
    rows = []
    for name, plaintext in inputs.items():
        for algorithm in ALGORITHMS:
            for kdf in ["pbkdf2", "scrypt", "argon2"]:
                try:
                    envelope = ck.encrypt_bytes(plaintext, PASSWORD, algorithm, kdf)
                    decrypted = ck.decrypt_bytes(envelope, PASSWORD)
                    ok = decrypted == plaintext
                    error = ""
                except Exception as e:  # pragma: no cover - dicatat sebagai kegagalan
                    ok = False
                    error = str(e)
                rows.append({
                    "input": name,
                    "ukuran_byte": len(plaintext),
                    "algoritma": algorithm,
                    "kdf": kdf,
                    "berhasil": ok,
                    "keterangan": error,
                })
    return rows


# ---------------------------------------------------------------------------
# 2) Waktu enkripsi & dekripsi untuk 1 KB, 1 MB, 10 MB
# ---------------------------------------------------------------------------

def run_timing_tests() -> list[dict]:
    sizes = {"1 KB": 1024, "1 MB": 1024 ** 2, "10 MB": 10 * 1024 ** 2}
    rows = []
    for size_label, size_bytes in sizes.items():
        plaintext = os.urandom(size_bytes)
        for algorithm in ALGORITHMS:
            enc_times, dec_times = [], []
            envelope = None
            for _ in range(TIMING_REPEATS):
                t0 = time.perf_counter()
                envelope = ck.encrypt_bytes(plaintext, PASSWORD, algorithm, KDF_FOR_SPEED_TESTS)
                enc_times.append(time.perf_counter() - t0)

                t0 = time.perf_counter()
                ck.decrypt_bytes(envelope, PASSWORD)
                dec_times.append(time.perf_counter() - t0)

            rows.append({
                "ukuran": size_label,
                "ukuran_byte": size_bytes,
                "algoritma": algorithm,
                "rata_rata_enkripsi_ms": round(statistics.mean(enc_times) * 1000, 3),
                "rata_rata_dekripsi_ms": round(statistics.mean(dec_times) * 1000, 3),
                "throughput_enkripsi_MBps": round((size_bytes / (1024 ** 2)) / statistics.mean(enc_times), 3),
                "throughput_dekripsi_MBps": round((size_bytes / (1024 ** 2)) / statistics.mean(dec_times), 3),
                "pengulangan": TIMING_REPEATS,
            })
    return rows


# ---------------------------------------------------------------------------
# 3) Avalanche effect
# ---------------------------------------------------------------------------

def run_avalanche_tests() -> list[dict]:
    sample_texts = [
        b"A",
        b"Pesan rahasia untuk pengujian avalanche effect.",
        os.urandom(256),
    ]
    rows = []
    for i, text in enumerate(sample_texts, start=1):
        for algorithm in ALGORITHMS:
            plain_pct = avalanche_effect_bit_flip(text, PASSWORD, algorithm, KDF_FOR_SPEED_TESTS)
            key_pct = avalanche_effect_key_change(text, PASSWORD, algorithm, KDF_FOR_SPEED_TESTS)
            rows.append({
                "sampel_ke": i,
                "panjang_plaintext_byte": len(text),
                "algoritma": algorithm,
                "avalanche_ubah_1_bit_plaintext_persen": plain_pct,
                "avalanche_ubah_1_bit_kunci_persen": key_pct,
                "ideal_acuan_persen": 50.0,
            })
    return rows


# ---------------------------------------------------------------------------
# 4) Entropi & histogram
# ---------------------------------------------------------------------------

def run_entropy_and_histogram(sample_plaintext: bytes) -> tuple[list[dict], list[int], list[int]]:
    rows = []
    hist_plain = byte_histogram(sample_plaintext)
    entropy_plain = shannon_entropy(sample_plaintext)
    rows.append({"jenis": "plaintext", "algoritma": "-", "entropi_bit_per_byte": entropy_plain})

    hist_cipher_last = hist_plain
    for algorithm in ALGORITHMS:
        envelope = ck.encrypt_bytes(sample_plaintext, PASSWORD, algorithm, KDF_FOR_SPEED_TESTS)
        entropy_cipher = shannon_entropy(envelope)
        rows.append({"jenis": "ciphertext", "algoritma": algorithm, "entropi_bit_per_byte": entropy_cipher})
        if algorithm == ALGORITHMS[-1]:
            hist_cipher_last = byte_histogram(envelope)

    return rows, hist_plain, hist_cipher_last


# ---------------------------------------------------------------------------
# 5) Perbandingan ringkas AES-256-GCM vs ChaCha20-Poly1305
# ---------------------------------------------------------------------------

def build_comparison_summary(timing_rows: list[dict], avalanche_rows: list[dict]) -> list[dict]:
    rows = []
    for algorithm in ALGORITHMS:
        timing_1mb = [r for r in timing_rows if r["algoritma"] == algorithm and r["ukuran"] == "1 MB"][0]
        avg_avalanche_plain = statistics.mean(
            r["avalanche_ubah_1_bit_plaintext_persen"] for r in avalanche_rows if r["algoritma"] == algorithm
        )
        rows.append({
            "algoritma": algorithm,
            "throughput_enkripsi_1MB_MBps": timing_1mb["throughput_enkripsi_MBps"],
            "throughput_dekripsi_1MB_MBps": timing_1mb["throughput_dekripsi_MBps"],
            "rata_rata_avalanche_plaintext_persen": round(avg_avalanche_plain, 3),
            "keterangan": (
                "AES-256-GCM: umumnya lebih cepat pada CPU dengan akselerasi AES-NI."
                if algorithm == "aes-gcm" else
                "ChaCha20-Poly1305: unggul pada perangkat tanpa akselerasi AES-NI "
                "(mis. mobile/embedded)."
            ),
        })
    return rows


# ---------------------------------------------------------------------------
# Penulisan hasil ke Excel (XLSX) + grafik histogram (PNG)
# ---------------------------------------------------------------------------

def write_excel_report(correctness, timing, avalanche, entropy, comparison):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def write_sheet(ws, rows, title):
        ws.title = title
        if not rows:
            return
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        for row in rows:
            ws.append([row[h] for h in headers])
        for i, header in enumerate(headers, start=1):
            width = max(12, min(40, len(str(header)) + 4))
            ws.column_dimensions[get_column_letter(i)].width = width

    write_sheet(wb.active, correctness, "Uji Korektnas")
    write_sheet(wb.create_sheet(), timing, "Uji Waktu")
    write_sheet(wb.create_sheet(), avalanche, "Avalanche Effect")
    write_sheet(wb.create_sheet(), entropy, "Entropi")
    write_sheet(wb.create_sheet(), comparison, "Perbandingan Algoritma")

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    out_path = os.path.join(OUTPUT_DIR, "hasil_pengujian.xlsx")
    wb.save(out_path)
    return out_path


def write_histogram_plots(hist_plain, hist_cipher):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    for name, hist in [("histogram_plaintext.png", hist_plain), ("histogram_ciphertext.png", hist_cipher)]:
        plt.figure(figsize=(8, 4))
        plt.bar(range(256), hist, width=1.0, color="#2F5496")
        plt.title(f"Histogram Byte - {name.replace('histogram_', '').replace('.png', '').capitalize()}")
        plt.xlabel("Nilai byte (0-255)")
        plt.ylabel("Frekuensi")
        plt.tight_layout()
        plt.savefig(os.path.join(OUTPUT_DIR, name), dpi=120)
        plt.close()


def main():
    print("[1/5] Menjalankan uji korektnas (>=10 masukan) ...")
    inputs = build_test_inputs()
    correctness = run_correctness_tests(inputs)
    n_ok = sum(1 for r in correctness if r["berhasil"])
    print(f"      {n_ok}/{len(correctness)} kombinasi input x algoritma x KDF berhasil.")

    print("[2/5] Mengukur waktu enkripsi/dekripsi (1 KB, 1 MB, 10 MB) ...")
    timing = run_timing_tests()

    print("[3/5] Menghitung avalanche effect ...")
    avalanche = run_avalanche_tests()

    print("[4/5] Menghitung entropi & histogram byte ...")
    sample_plaintext = inputs["teks_panjang_lorem"]
    entropy, hist_plain, hist_cipher = run_entropy_and_histogram(sample_plaintext)

    print("[5/5] Menyusun ringkasan perbandingan algoritma ...")
    comparison = build_comparison_summary(timing, avalanche)

    xlsx_path = write_excel_report(correctness, timing, avalanche, entropy, comparison)
    write_histogram_plots(hist_plain, hist_cipher)

    print("\nSelesai.")
    print(f"  - Tabel hasil pengujian : {xlsx_path}")
    print(f"  - Grafik histogram      : {OUTPUT_DIR}/histogram_plaintext.png, histogram_ciphertext.png")
    print("\nSalin nilai-nilai ini ke Bagian 5 (Pengujian dan Analisis) pada laporan teknis.")


if __name__ == "__main__":
    main()
